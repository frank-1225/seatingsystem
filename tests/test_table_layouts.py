import io
import unittest

from openpyxl import load_workbook

from app import app


class TableLayoutExportTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()

    @staticmethod
    def attendee(name):
        return {
            "id": 1,
            "name": name,
            "unit": "测试单位",
            "category": "嘉宾",
            "phone": "13800138000",
            "extra": {},
        }

    def export_workbook(self, payload):
        response = self.client.post("/api/export", json=payload)
        self.assertEqual(response.status_code, 200, response.get_json(silent=True))
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return load_workbook(io.BytesIO(response.data))

    def test_index_exposes_long_round_and_u_table_modes(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("长会议桌式（对称围坐）", html)
        self.assertIn("圆桌式（围坐）", html)
        self.assertIn("U型桌式（领导席＋两侧参会席）", html)
        self.assertIn("主宾双主位", html)
        self.assertIn('class="classroom-layout"', html)
        self.assertIn('v-model.number="sideCfg.seatGap"', html)
        self.assertIn('v-model.number="sideCfg.seatsPerRow"', html)
        self.assertIn('min="0" max="20"', html)
        self.assertIn("longtableSideSpan('top')", html)
        self.assertIn("longtableMinimumSpan", html)

    def test_long_table_export_uses_human_readable_seat_names(self):
        payload = {
            "conf_name": "长桌测试会",
            "layoutType": "longtable",
            "podium": {},
            "main": {},
            "longtableConfig": {
                "sides": {
                    "top": {"label": "北侧"},
                    "bottom": {"label": "南侧"},
                    "left": {"label": "西侧"},
                    "right": {"label": "东侧"},
                }
            },
            "roundtableConfig": {},
            "seats_data": {
                "lt-top-0-0": {
                    "type": "normal",
                    "attendee": self.attendee("长桌嘉宾"),
                    "region_name": "未分区",
                }
            },
            "sms_template": "{姓名}：{排号} / {座号}",
            "export_template": "{姓名}",
            "region_types": [],
        }

        workbook = self.export_workbook(payload)
        self.assertEqual(workbook["会场布局图"]["A1"].value, "长桌测试会 - 长会议桌式会议")

        sign_in_rows = list(workbook["参会人员签到表"].iter_rows(values_only=True))
        self.assertIn(
            ("长会议桌式会议", "未分区", "北侧", "第1排 第1座", "长桌嘉宾", "测试单位", "嘉宾", "13800138000", None),
            sign_in_rows,
        )

        sms_rows = list(workbook["短信群发清单"].iter_rows(values_only=True))
        self.assertIn(
            ("长桌嘉宾", "测试单位", "13800138000", "长桌嘉宾：北侧 / 第1排 第1座"),
            sms_rows,
        )

    def test_round_table_export_includes_table_and_ring(self):
        payload = {
            "conf_name": "圆桌测试会",
            "layoutType": "roundtable",
            "podium": {},
            "main": {},
            "longtableConfig": {},
            "roundtableConfig": {
                "tables": 2,
                "rings": [{"id": "ring-0", "label": "主桌", "seatCount": 10, "radius": 150}],
            },
            "seats_data": {
                "rt-1-ring-0-2": {
                    "type": "normal",
                    "attendee": self.attendee("圆桌嘉宾"),
                    "region_name": "未分区",
                }
            },
            "sms_template": "{姓名}：{排号} / {座号}",
            "export_template": "{姓名}",
            "region_types": [],
        }

        workbook = self.export_workbook(payload)
        self.assertEqual(workbook["会场布局图"]["A1"].value, "圆桌测试会 - 圆桌式会议")

        sign_in_rows = list(workbook["参会人员签到表"].iter_rows(values_only=True))
        self.assertIn(
            ("圆桌式会议", "未分区", "第2桌-主桌", "第3座", "圆桌嘉宾", "测试单位", "嘉宾", "13800138000", None),
            sign_in_rows,
        )

    def test_round_table_dual_principal_export_names_both_principal_seats(self):
        payload = {
            "conf_name": "双主位测试会",
            "layoutType": "roundtable",
            "podium": {},
            "main": {},
            "longtableConfig": {},
            "roundtableConfig": {
                "tables": 1,
                "dualPrincipal": True,
                "rings": [{"id": "ring-0", "label": "主桌", "seatCount": 10, "radius": 150}],
            },
            "seats_data": {
                "rt-0-ring-0-0": {"type": "normal", "attendee": self.attendee("主人一号"), "region_name": "未分区"},
                "rt-0-ring-0-1": {"type": "normal", "attendee": self.attendee("宾客一号"), "region_name": "未分区"},
            },
            "sms_template": "{姓名}：{座号}",
            "export_template": "{姓名}",
            "region_types": [],
        }

        workbook = self.export_workbook(payload)
        sign_in_rows = list(workbook["参会人员签到表"].iter_rows(values_only=True))
        self.assertTrue(any(row[4] == "主人一号" and row[3] == "主人主位" for row in sign_in_rows))
        self.assertTrue(any(row[4] == "宾客一号" and row[3] == "宾客主位" for row in sign_in_rows))

    def test_u_table_export_describes_leader_and_near_side_positions(self):
        payload = {
            "conf_name": "U型桌测试会",
            "layoutType": "utable",
            "podium": {},
            "main": {},
            "longtableConfig": {},
            "roundtableConfig": {},
            "utableConfig": {"leaderSeatCount": 5, "leaderSide": "top"},
            "seats_data": {
                "u-leader-2": {"type": "normal", "attendee": self.attendee("领导一号"), "region_name": "未分区"},
                "u-left-0": {"type": "normal", "attendee": self.attendee("参会一号"), "region_name": "未分区"},
            },
            "sms_template": "{姓名}：{排号} / {座号}",
            "export_template": "{姓名}",
            "region_types": [],
        }

        workbook = self.export_workbook(payload)
        self.assertEqual(workbook["会场布局图"]["A1"].value, "U型桌测试会 - U型桌式会议")
        sign_in_rows = list(workbook["参会人员签到表"].iter_rows(values_only=True))
        self.assertTrue(any(row[4] == "领导一号" and row[2:4] == ("领导席", "第1位") for row in sign_in_rows))
        self.assertTrue(any(row[4] == "参会一号" and row[2:4] == ("左侧参会席", "距领导第1位") for row in sign_in_rows))

    def test_u_table_export_includes_side_column_for_multicolumn_layout(self):
        payload = {
            "conf_name": "U型桌多列测试会",
            "layoutType": "utable",
            "podium": {},
            "main": {},
            "longtableConfig": {},
            "roundtableConfig": {},
            "utableConfig": {"leaderSeatCount": 5, "rightColumnCount": 2, "rightSeatCount": 4},
            "seats_data": {
                "u-right-1-2": {"type": "normal", "attendee": self.attendee("外列参会人"), "region_name": "未分区"},
            },
            "sms_template": "{姓名}：{排号} / {座号}",
            "export_template": "{姓名}",
            "region_types": [],
        }

        workbook = self.export_workbook(payload)
        sign_in_rows = list(workbook["参会人员签到表"].iter_rows(values_only=True))
        self.assertTrue(any(row[4] == "外列参会人" and row[2:4] == ("右侧参会席-第2列", "距领导第3位") for row in sign_in_rows))


if __name__ == "__main__":
    unittest.main()
